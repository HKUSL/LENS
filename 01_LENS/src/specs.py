from __future__ import annotations

import os
from typing import Dict, Optional, Tuple

import docx
from docx.oxml.ns import qn

from .common import DEFAULT_SPECS_DIR


INTERFACE_SPEC_MAP = {
    "s6a": "29272-j30.docx",
    "s6d": "29272-j30.docx",
    "cx": "29228-j10.docx",
    "sh": "29329-j30.docx",
    "n1": "24501-h90.docx",
    "sip": "RFC 3261.txt",
    "mqtt": "mqttv3.11.txt",
    "coap": "RFC 7252.txt",
    "oci": "oci-runtime-spec.txt",
    "sle": "ccsds-sle.txt",
    "tc": "ccsds-tc-sdls.txt",
    "wayland": "wayland-protocol.txt",
}


class SpecCache:
    def __init__(self, specs_dir: str = DEFAULT_SPECS_DIR):
        self.specs_dir = specs_dir
        self._cache: Dict[str, str] = {}

    def get_for_interface(self, interface: str) -> Optional[Tuple[str, str]]:
        name = INTERFACE_SPEC_MAP.get(interface.lower())
        if not name:
            return None
        path = os.path.join(self.specs_dir, name)
        if not os.path.exists(path):
            return None
        if name not in self._cache:
            self._cache[name] = get_full_text(path)
        return name, self._cache[name]


def get_full_text(path: str) -> str:
    if path.endswith(".txt"):
        with open(path, "r", encoding="utf-8") as f:
            return f.read()

    doc = docx.Document(path)
    paragraphs = {p._element: p for p in doc.paragraphs}
    tables = {t._element: t for t in doc.tables}
    parts = []

    for child in doc.element.body:
        if child.tag == qn("w:p") and child in paragraphs:
            text = paragraphs[child].text.strip()
            if text:
                parts.append(text)
        elif child.tag == qn("w:tbl") and child in tables:
            parts.append(_table_to_text(tables[child]))
    return "\n".join(parts)


def _table_to_text(table) -> str:
    rows = []
    for row in table.rows:
        rows.append(" | ".join(cell.text.strip().replace("\n", " ") for cell in row.cells))
    return "\n".join(rows)

